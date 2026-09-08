import os
import glob
import json
import pandas as pd
import openpyxl

# Rutas de datos
DATA_DIR = "data"
CHEQUES_DIR = os.path.join(DATA_DIR, "cheques")
OUTPUT_HTML = "index.html"

def load_saldos():
    """Procesa el Archivo 1: Saldos Operativos"""
    path = os.path.join(DATA_DIR, "saldos_operativos.xlsx")
    if not os.path.exists(path):
        print(f"Advertencia: No se encontró {path}")
        return []
    
    # Carga columnas M (Empresa), N (Banco), O (Saldo Inicial), P (Saldo Operativo)
    df = pd.read_excel(path, usecols="M:P", header=0)
    df.columns = ["empresa", "banco", "saldo_inicial", "saldo_operativo"]
    
    # Limpieza
    df = df.dropna(subset=["empresa", "banco"])
    df["empresa"] = df["empresa"].astype(str).str.strip()
    df["banco"] = df["banco"].astype(str).str.strip()
    df["saldo_inicial"] = pd.to_numeric(df["saldo_inicial"], errors="coerce").fillna(0)
    df["saldo_operativo"] = pd.to_numeric(df["saldo_operativo"], errors="coerce").fillna(0)
    
    # Por defecto, asumimos ARS si no especifica moneda en este archivo
    df["moneda"] = "ARS"
    return df.to_dict(orient="records")

def load_inversiones():
    """Procesa el Archivo 2: Inversiones (Pestaña 'CARGA')"""
    path = os.path.join(DATA_DIR, "inversiones.xlsx")
    if not os.path.exists(path):
        print(f"Advertencia: No se encontró {path}")
        return []
    
    # Pestaña CARGA, cols B a K
    df = pd.read_excel(path, sheet_name="CARGA", usecols="B:K", header=0)
    df.columns = [
        "empresa", "tipo_inversion", "movimiento", "fondo", 
        "capital", "banco", "liquidez", "moneda", "vencimiento", "tasa"
    ]
    
    df = df.dropna(subset=["empresa", "capital"])
    df["empresa"] = df["empresa"].astype(str).str.strip()
    df["movimiento"] = df["movimiento"].astype(str).str.upper().str.strip()
    df["capital"] = pd.to_numeric(df["capital"], errors="coerce").fillna(0)
    df["moneda"] = df["moneda"].astype(str).str.strip().str.upper()
    df["liquidez"] = df["liquidez"].astype(str).str.strip()
    
    # Calcular monto neto: Rescate resta, Suscripción suma
    df["monto_neto"] = df.apply(
        lambda r: -r["capital"] if "RESCATE" in r["movimiento"] else r["capital"], axis=1
    )
    
    return df.to_dict(orient="records")

def load_cheques():
    """Procesa el Archivo 3: Archivos individuales por empresa en data/cheques/"""
    cheques_data = []
    files = glob.glob(os.path.join(CHEQUES_DIR, "*.xlsx"))
    
    for file_path in files:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            if "resumen viejo" not in wb.sheetnames:
                continue
            
            ws = wb["resumen viejo"]
            
            # Buscar Empresa en A1 o B1
            empresa_raw = str(ws["A1"].value or ws["B1"].value or "")
            empresa = empresa_raw.replace("EMPRESA:", "").strip()
            
            # Detectar Bancos desde columna D en la Fila 1
            bancos = []
            col_idx = 4 # Columna D
            while True:
                banco_val = ws.cell(row=1, column=col_idx).value
                if banco_val and str(banco_val).strip() != "-" and str(banco_val).strip().lower() != "totales":
                    bancos.append((col_idx, str(banco_val).strip()))
                    col_idx += 1
                else:
                    break
            
            # Recorrer filas de conceptos
            for row in range(2, 10):
                concepto = str(ws.cell(row=row, column=1).value or ws.cell(row=row, column=2).value or "").strip()
                if not concepto or "Totales" in concepto:
                    continue
                
                for c_idx, b_nombre in bancos:
                    monto = ws.cell(row=row, column=c_idx).value
                    try:
                        monto = float(monto) if monto is not None else 0.0
                    except ValueError:
                        monto = 0.0
                    
                    if monto > 0:
                        cheques_data.append({
                            "empresa": empresa,
                            "banco": b_nombre,
                            "concepto": concepto,
                            "monto": monto,
                            "moneda": "ARS" # Moneda por defecto
                        })
        except Exception as e:
            print(f"Error procesando cheques en {file_path}: {e}")
            
    return cheques_data

def generate_html(saldos, inversiones, cheques):
    """Genera el HTML estático interactivo con Tailwind CSS y Chart.js"""
    saldos_json = json.dumps(saldos)
    inversiones_json = json.dumps(inversiones)
    cheques_json = json.dumps(cheques)
    
    html_template = f"""<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Tablero de Tesorería Consolidada</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
</head>
<body class="bg-slate-900 text-slate-100 font-sans min-h-screen p-6">
    
    <!-- Encabezado y Filtros -->
    <header class="flex flex-col md:flex-row justify-between items-center pb-6 mb-6 border-b border-slate-700 gap-4">
        <div>
            <h1 class="text-2xl font-bold text-blue-400">Posición Diaria de Tesorería</h1>
            <p class="text-xs text-slate-400">Consolidación en tiempo real por grupo de empresas</p>
        </div>
        <div class="flex items-center gap-4">
            <!-- Selector Moneda -->
            <div class="flex bg-slate-800 rounded-lg p-1 border border-slate-700">
                <button id="btnARS" onclick="setMoneda('ARS')" class="px-3 py-1 text-xs font-semibold rounded-md bg-blue-600 text-white">ARS ($)</button>
                <button id="btnUSD" onclick="setMoneda('USD')" class="px-3 py-1 text-xs font-semibold rounded-md text-slate-400 hover:text-white">USD ($)</button>
            </div>
            <!-- Selector Empresa -->
            <select id="empresaSelect" onchange="renderDashboard()" class="bg-slate-800 border border-slate-700 rounded-lg text-xs px-3 py-2 text-slate-200">
                <option value="TODAS">-- Todas las Empresas --</option>
            </select>
        </div>
    </header>

    <!-- Tarjetas de Resumen (KPIs) -->
    <div class="grid grid-cols-1 md:grid-cols-4 gap-4 mb-6">
        <div class="bg-slate-800 p-4 rounded-xl border border-slate-700">
            <p class="text-xs text-slate-400 uppercase font-semibold">Saldo Operativo Bancos</p>
            <h2 id="kpiSaldos" class="text-xl font-bold text-emerald-400 mt-1">$ 0,00</h2>
        </div>
        <div class="bg-slate-800 p-4 rounded-xl border border-slate-700">
            <p class="text-xs text-slate-400 uppercase font-semibold">Total Inversiones</p>
            <h2 id="kpiInversiones" class="text-xl font-bold text-blue-400 mt-1">$ 0,00</h2>
        </div>
        <div class="bg-slate-800 p-4 rounded-xl border border-slate-700">
            <p class="text-xs text-slate-400 uppercase font-semibold">Cheques Pendientes</p>
            <h2 id="kpiCheques" class="text-xl font-bold text-rose-400 mt-1">$ 0,00</h2>
        </div>
        <div class="bg-slate-800 p-4 rounded-xl border border-slate-700">
            <p class="text-xs text-slate-400 uppercase font-semibold">Posición Liquida Neta</p>
            <h2 id="kpiPosicionNeta" class="text-xl font-bold text-cyan-300 mt-1">$ 0,00</h2>
        </div>
    </div>

    <!-- Gráficos y Tablas -->
    <div class="grid grid-cols-1 lg:grid-cols-2 gap-6 mb-6">
        <div class="bg-slate-800 p-4 rounded-xl border border-slate-700">
            <h3 class="text-sm font-semibold text-slate-300 mb-3">Saldos por Banco</h3>
            <canvas id="chartBancos" height="180"></canvas>
        </div>
        <div class="bg-slate-800 p-4 rounded-xl border border-slate-700">
            <h3 class="text-sm font-semibold text-slate-300 mb-3">Inversiones por Perfil de Liquidez</h3>
            <canvas id="chartLiquidez" height="180"></canvas>
        </div>
    </div>

    <!-- Cargar Datos JS -->
    <script>
        const saldosData = {saldos_json};
        const inversionesData = {inversiones_json};
        const chequesData = {cheques_json};
        
        let monedaActual = 'ARS';
        let chartBancosInst = null;
        let chartLiquidezInst = null;

        function setMoneda(m) {{
            monedaActual = m;
            document.getElementById('btnARS').className = m === 'ARS' ? "px-3 py-1 text-xs font-semibold rounded-md bg-blue-600 text-white" : "px-3 py-1 text-xs font-semibold rounded-md text-slate-400 hover:text-white";
            document.getElementById('btnUSD').className = m === 'USD' ? "px-3 py-1 text-xs font-semibold rounded-md bg-blue-600 text-white" : "px-3 py-1 text-xs font-semibold rounded-md text-slate-400 hover:text-white";
            renderDashboard();
        }}

        function initEmpresaFilter() {{
            const empresas = new Set();
            saldosData.forEach(d => empresas.add(d.empresa));
            inversionesData.forEach(d => empresas.add(d.empresa));
            chequesData.forEach(d => empresas.add(d.empresa));

            const select = document.getElementById('empresaSelect');
            empresas.forEach(e => {{
                if(e && e !== 'nan') {{
                    const opt = document.createElement('option');
                    opt.value = e;
                    opt.textContent = e;
                    select.appendChild(opt);
                }}
            }});
        }}

        function fmt(val) {{
            return new Intl.NumberFormat('es-AR', {{ style: 'currency', currency: monedaActual }}).format(val);
        }}

        function renderDashboard() {{
            const emp = document.getElementById('empresaSelect').value;

            // Filtrar Saldos
            const fSaldos = saldosData.filter(d => (emp === 'TODAS' || d.empresa === emp) && d.moneda === monedaActual);
            const totalSaldos = fSaldos.reduce((acc, d) => acc + d.saldo_operativo, 0);

            // Filtrar Inversiones
            const fInv = inversionesData.filter(d => (emp === 'TODAS' || d.empresa === emp) && d.moneda === monedaActual);
            const totalInv = fInv.reduce((acc, d) => acc + d.monto_neto, 0);

            // Filtrar Cheques
            const fCheques = chequesData.filter(d => (emp === 'TODAS' || d.empresa === emp) && d.moneda === monedaActual);
            const totalCheques = fCheques.reduce((acc, d) => acc + d.monto, 0);

            // Actualizar KPIs
            document.getElementById('kpiSaldos').textContent = fmt(totalSaldos);
            document.getElementById('kpiInversiones').textContent = fmt(totalInv);
            document.getElementById('kpiCheques').textContent = fmt(totalCheques);
            document.getElementById('kpiPosicionNeta').textContent = fmt(totalSaldos + totalInv - totalCheques);

            // Chart Bancos
            const bancosMap = {{}};
            fSaldos.forEach(d => {{ bancosMap[d.banco] = (bancosMap[d.banco] || 0) + d.saldo_operativo; }});
            
            if(chartBancosInst) chartBancosInst.destroy();
            chartBancosInst = new Chart(document.getElementById('chartBancos'), {{
                type: 'bar',
                data: {{
                    labels: Object.keys(bancosMap),
                    datasets: [{{ label: 'Saldo Operativo', data: Object.values(bancosMap), backgroundColor: '#3b82f6' }}]
                }},
                options: {{ plugins: {{ legend: {{ display: false }} }}, scales: {{ y: {{ ticks: {{ color: '#94a3b8' }} }}, x: {{ ticks: {{ color: '#94a3b8' }} }} }} }}
            }});

            // Chart Liquidez
            const liqMap = {{}};
            fInv.forEach(d => {{ liqMap[d.liquidez] = (liqMap[d.liquidez] || 0) + d.monto_neto; }});
            
            if(chartLiquidezInst) chartLiquidezInst.destroy();
            chartLiquidezInst = new Chart(document.getElementById('chartLiquidez'), {{
                type: 'doughnut',
                data: {{
                    labels: Object.keys(liqMap),
                    datasets: [{{ data: Object.values(liqMap), backgroundColor: ['#10b981', '#f59e0b', '#6366f1', '#ec4899'] }}]
                }},
                options: {{ plugins: {{ legend: {{ labels: {{ color: '#94a3b8' }} }} }} }}
            }});
        }}

        window.onload = () => {{
            initEmpresaFilter();
            renderDashboard();
        }};
    </script>
</body>
</html>
"""
    with open(OUTPUT_HTML, "w", encoding="utf-8") as f:
        f.write(html_template)
    print(f"Dashboard generado exitosamente en {OUTPUT_HTML}")

if __name__ == "__main__":
    saldos = load_saldos()
    inversiones = load_inversiones()
    cheques = load_cheques()
    generate_html(saldos, inversiones, cheques)
